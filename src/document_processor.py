import os
from pathlib import Path
from typing import List, Dict, Any
import pypdf
from docx import Document as DocxDocument
import markdown
from bs4 import BeautifulSoup
from striprtf.striprtf import rtf_to_text
from config import config


class DocumentProcessor:
    def __init__(self):
        self.supported_extensions = {'.pdf', '.docx', '.txt', '.md', '.html', '.htm', '.rtf'}

    def extract_text(self, file_path: Path) -> str:
        """Extract text from various document formats."""
        ext = file_path.suffix.lower()

        try:
            if ext == '.pdf':
                return self._extract_pdf(file_path)
            elif ext == '.docx':
                return self._extract_docx(file_path)
            elif ext in ['.txt', '.md']:
                return self._extract_text_file(file_path)
            elif ext in ['.html', '.htm']:
                return self._extract_html(file_path)
            elif ext == '.rtf':
                return self._extract_rtf(file_path)
            else:
                raise ValueError(f"Unsupported file format: {ext}")
        except Exception as e:
            raise RuntimeError(f"Failed to extract text from {file_path}: {e}")

    def _extract_pdf(self, file_path: Path) -> str:
        text_parts = []
        with open(file_path, 'rb') as f:
            reader = pypdf.PdfReader(f)
            for page in reader.pages:
                text = page.extract_text()
                if text:
                    text_parts.append(text)
        return '\n\n'.join(text_parts)

    def _extract_docx(self, file_path: Path) -> str:
        doc = DocxDocument(file_path)
        # Join paragraphs and also pull tables, which doc.paragraphs omits.
        parts: List[str] = []
        for para in doc.paragraphs:
            if para.text.strip():
                parts.append(para.text)
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                line = " | ".join(c for c in cells if c)
                if line:
                    parts.append(line)
        return '\n\n'.join(parts)

    def _extract_text_file(self, file_path: Path) -> str:
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            return f.read()

    def _extract_html(self, file_path: Path) -> str:
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            soup = BeautifulSoup(f.read(), 'html.parser')
            return soup.get_text(separator='\n', strip=True)

    def _extract_rtf(self, file_path: Path) -> str:
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            return rtf_to_text(f.read())

    def chunk_text(self, text: str, chunk_size: int = None, overlap: int = None) -> List[str]:
        """Split text into overlapping chunks."""
        chunk_size = chunk_size or config.chunk_size
        overlap = overlap or config.chunk_overlap

        if len(text) <= chunk_size:
            return [text]

        chunks = []
        start = 0
        while start < len(text):
            end = start + chunk_size
            chunk = text[start:end]

            # Try to break at a sentence boundary
            if end < len(text):
                last_period = chunk.rfind('. ')
                last_newline = chunk.rfind('\n')
                break_point = max(last_period, last_newline)
                if break_point > chunk_size // 2:
                    chunk = chunk[:break_point + 1]
                    end = start + break_point + 1

            chunks.append(chunk.strip())
            start = end - overlap

        return [c for c in chunks if c.strip()]

    @staticmethod
    def source_key(file_path: Path) -> str:
        """Return a stable, location-independent key for a document.

        The key is the path relative to the project root (or the bare
        filename if the file lives outside the project). This means the
        same document produces the same key regardless of which machine or
        absolute path it was ingested from, so skip-detection, summarizing
        and lookups work consistently.
        """
        try:
            return str(Path(file_path).resolve().relative_to(config.project_root.resolve()))
        except ValueError:
            return Path(file_path).name

    @staticmethod
    def match_source(query: str) -> str:
        """Resolve a user-provided source argument to a stored 'source' key.

        Accepts an existing path, an absolute path, a basename, or a path
        relative to the documents directory, and returns the canonical key
        that is stored in the vector database.
        """
        q = str(query).strip().strip('"').strip("'")
        p = Path(q)
        if p.exists() or p.is_absolute():
            return DocumentProcessor.source_key(p)
        # Otherwise treat it as a path under the documents directory.
        return DocumentProcessor.source_key(Path(config.documents_path) / q)

    def process_file(self, file_path: Path) -> List[Dict[str, Any]]:
        """Process a single file into chunks with metadata."""
        text = self.extract_text(file_path)
        chunks = self.chunk_text(text)

        # Stable, location-independent source key so lookups work regardless of
        # the absolute path used to ingest the file.
        source_key = self.source_key(file_path)
        abs_path = str(Path(file_path).resolve())

        return [
            {
                'text': chunk,
                'source': source_key,        # stable relative key (the lookup key)
                'source_path': abs_path,      # original absolute path (reference only)
                'filename': file_path.name,
                'chunk_index': i,
                'total_chunks': len(chunks)
            }
            for i, chunk in enumerate(chunks)
        ]

    def scan_documents(self) -> List[Path]:
        """Scan documents directory for supported files."""
        doc_dir = Path(config.documents_path)
        if not doc_dir.exists():
            return []

        files = []
        for ext in self.supported_extensions:
            files.extend(doc_dir.rglob(f'*{ext}'))
        # De-duplicate while preserving a stable order (rglob can repeat).
        seen = set()
        unique = []
        for f in sorted(files, key=lambda p: str(p)):
            rp = f.resolve()
            if rp not in seen:
                seen.add(rp)
                unique.append(f)
        return unique


def main():
    processor = DocumentProcessor()
    files = processor.scan_documents()
    print(f"Found {len(files)} documents:")
    for f in files:
        print(f"    - {f}")


if __name__ == '__main__':
    main()
